"""Build Estudio_Global_Metricas_3D.xlsx from exports/dataset3d_metrics.csv."""
import argparse

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

M = ['area_mm2', 'perimeter_mm', 'major_axis_length_mm', 'minor_axis_length_mm',
     'depth_mean_mm', 'depth_p95_mm', 'depth_max_mm', 'depth_volume_mm3']
SHEETS = {'Holes': 'Material loss - Holes', 'Crack': 'Crack (cut)',
          'Friction': 'Scratch-Friction', 'Abrasion': 'Scratch-Abrasion'}


COLS = ['relative_path', 'defect_type', 'variant', 'fabric', 'detected', 'pred_class', 'conf',
        'area_mm2', 'perimeter_mm', 'major_axis_length_mm', 'minor_axis_length_mm',
        'eccentricity', 'compactation', 'glcm_homogeneity', 'glcm_energy', 'glcm_entropy',
        'depth_mean_mm', 'depth_p95_mm', 'depth_max_mm', 'depth_volume_mm3']
SIZE = {'Holes': 'area_mm2', 'Abrasion': 'area_mm2', 'Friction': 'area_mm2',
        'Crack': 'major_axis_length_mm'}
THICK = {'Holes': 'multiple', 'Abrasion': 'thick', 'Friction': 'thick', 'Crack': 'thick'}
QUANTILES = (0.1, 0.25, 0.5, 0.75, 0.9)


def col_letter(name):
    return get_column_letter(COLS.index(name) + 1)


class Blocks:
    """Row ranges of each class sheet, sorted so every variant/fabric is contiguous."""

    def __init__(self, det):
        self.data, self.rng = {}, {}
        for kind, g in det.groupby('defect_type'):
            g = g.sort_values(['variant', 'fabric', 'relative_path']).reset_index(drop=True)
            self.data[kind] = g
            name = SHEETS[kind]
            for key, idx in {**g.groupby('variant').groups,
                             **{(v, f): i for (v, f), i in g.groupby(['variant', 'fabric']).groups.items()},
                             'all': g.index}.items():
                self.rng[(kind, key)] = (name, idx.min() + 2, idx.max() + 2)

    def ref(self, kind, key, metric):
        name, r1, r2 = self.rng[(kind, key)]
        c = col_letter(metric)
        return f"'{name}'!${c}${r1}:${c}${r2}"


def header(ws, row, values):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDE6F0')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--csv', default='exports/dataset3d_metrics.csv')
    ap.add_argument('--only', default='exports/keep_2d_images.txt',
                    help='relative paths kept (images present in the 2D study)')
    ap.add_argument('--out', default='../Estudio_Global_Metricas_3D.xlsx')
    a = ap.parse_args()
    df = pd.read_csv(a.csv)
    if a.only:
        df = df[df.relative_path.isin(set(open(a.only).read().split()))]
    det = df[df.detected == 1].copy()
    images = df.groupby(['defect_type', 'variant']).size()
    blocks = Blocks(det)

    note = pd.DataFrame({'Notas': [
        'Profundidad estimada por la red RGB->normales + Poisson (gs3drecon), no medida.',
        'Referencia de cero: mediana de las imagenes without_defect de cada tela.',
        'Profundidad en mm relativa al plano local del anillo alrededor de la mascara.',
        'Se desconoce la fuerza de contacto de este dataset.',
        'Solo imagenes presentes en Estudio_Global_Metricas.xlsx (2D, sin las imagenes malas): thin y hole fabric01.',
        'Una region por imagen (la de mayor confianza de YOLO12 distinta de No Defect).',
        'Holes: fabric01 = simple, fabric02-06 = multiple. Abrasion/Friction/Crack: thin/thick.',
        'Los umbrales de reparabilidad deben salir de un criterio real, no de estas tablas.']})

    with pd.ExcelWriter(a.out, engine='openpyxl') as w:
        wb = w.book
        wb.remove(wb.active) if wb.worksheets else None
        order = ['Rangos Globales 3D', 'Por tela', 'Umbrales propuestos',
                 'Percentiles profundidad', 'Cobertura deteccion']
        for name in order:
            wb.create_sheet(name)
        for kind in SHEETS:
            blocks.data[kind][COLS].to_excel(w, sheet_name=SHEETS[kind], index=False)
        note.to_excel(w, sheet_name='Notas', index=False)

        def stat_rows(ws, keys, label_cols):
            head = label_cols + ['n_detected']
            for m in M:
                head += [f'{m}_min', f'{m}_max', f'{m}_mean']
            header(ws, 1, head)
            r = 2
            for kind, key, labels in keys:
                for i, v in enumerate(labels, 1):
                    ws.cell(row=r, column=i, value=v)
                c = len(labels) + 1
                ws.cell(row=r, column=c, value=f"=COUNT({blocks.ref(kind, key, 'area_mm2')})")
                for m in M:
                    ref = blocks.ref(kind, key, m)
                    for f in ('MIN', 'MAX', 'AVERAGE'):
                        c += 1
                        ws.cell(row=r, column=c, value=f'={f}({ref})')
                r += 1

        keys, fab = [], []
        for kind, g in blocks.data.items():
            for v in sorted(g.variant.unique()):
                keys.append((kind, v, [kind, v]))
                for f in sorted(g[g.variant == v].fabric.unique()):
                    fab.append((kind, (v, f), [kind, v, f]))
        stat_rows(w.sheets['Rangos Globales 3D'], keys, ['defect_type', 'variant'])
        stat_rows(w.sheets['Por tela'], fab, ['defect_type', 'variant', 'fabric'])

        ws = w.sheets['Percentiles profundidad']
        qm = ['depth_p95_mm', 'depth_max_mm', 'depth_volume_mm3', 'area_mm2']
        header(ws, 1, ['defect_type', 'variant'] + [f'{m}_P{int(q*100)}' for m in qm for q in QUANTILES])
        r = 2
        for kind, key, labels in keys:
            for i, v in enumerate(labels, 1):
                ws.cell(row=r, column=i, value=v)
            c = 2
            for m in qm:
                for q in QUANTILES:
                    c += 1
                    ws.cell(row=r, column=c, value=f'=PERCENTILE({blocks.ref(kind, key, m)},{q})')
            r += 1

        ws = w.sheets['Umbrales propuestos']
        header(ws, 1, ['defect_type', 'criterio', 'percentil', 'umbral_depth_p95_mm',
                       'metrica_tamano', 'umbral_tamano', 'umbral_volumen_mm3',
                       '%_no_reparable', 'regla'])
        r = 2
        for kind in blocks.data:
            size = SIZE[kind]
            rows = [(f'P{int(q*100)} de todos', 'all', q) for q in (0.5, 0.75, 0.9)]
            if (kind, THICK[kind]) in blocks.rng:
                rows.append((f'P10 de {THICK[kind]}', THICK[kind], 0.1))
            for label, key, q in rows:
                ws.cell(row=r, column=1, value=kind)
                ws.cell(row=r, column=2, value=label)
                ws.cell(row=r, column=3, value=q)
                for col, metric in ((4, 'depth_p95_mm'), (6, size), (7, 'depth_volume_mm3')):
                    ws.cell(row=r, column=col, value=f'=PERCENTILE({blocks.ref(kind, key, metric)},$C{r})')
                ws.cell(row=r, column=5, value=size)
                d, s_, v_ = (blocks.ref(kind, 'all', m) for m in ('depth_p95_mm', size, 'depth_volume_mm3'))
                ws.cell(row=r, column=8, value=(
                    f'=SUMPRODUCT(--((({d}>$D{r})+({s_}>$F{r})+({v_}>$G{r}))>0))/COUNT({d})'))
                ws.cell(row=r, column=8).number_format = '0.0%'
                ws.cell(row=r, column=9, value='NO reparable si depth_p95 > umbral O tamano > umbral O volumen > umbral')
                r += 1

        ws = w.sheets['Cobertura deteccion']
        header(ws, 1, ['defect_type', 'variant', 'images', 'detected', 'detection_rate'])
        r = 2
        for kind, key, labels in keys:
            ws.cell(row=r, column=1, value=kind)
            ws.cell(row=r, column=2, value=key)
            ws.cell(row=r, column=3, value=int(images[(kind, key)]))
            ws.cell(row=r, column=4, value=f"=COUNT({blocks.ref(kind, key, 'area_mm2')})")
            ws.cell(row=r, column=5, value=f'=D{r}/C{r}')
            ws.cell(row=r, column=5).number_format = '0.0%'
            r += 1

        for ws in wb.worksheets:
            ws.freeze_panes = 'A2'
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    38, max(10, len(str(col[0].value)) + 2))
    print('written', a.out, len(det), 'regions')


if __name__ == '__main__':
    main()
